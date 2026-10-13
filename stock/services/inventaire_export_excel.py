"""
Export Excel du rapport d'un inventaire validé.

Le classeur reprend la présentation du rapport d'inventaire (en-tête de l'établissement,
informations de la session, indicateurs, tableau des articles, totaux et écarts financiers)
avec les couleurs configurées pour le rapport « inventaire » de l'entreprise.

Seules deux colonnes sont modifiables : « Stock physique » et « Dernier PU ». Tout le reste
est verrouillé (feuille protégée) et les montants, écarts, indicateurs et totaux sont de
vraies formules Excel, identiques aux règles du backend :

- Écart = Stock physique − Stock logiciel                     (InventaireLigne.recalculer_ecart)
- Montant logiciel = Stock logiciel × Dernier PU               (InventaireLigne.montant_logiciel)
- Montant physique = Stock physique × Dernier PU               (InventaireLigne.montant_physique)
- Écart (montant) = Montant physique − Montant logiciel = Écart × Dernier PU  (ecart_montant)
- Total logiciel / physique = sommes des montants ; Écart financier = physique − logiciel
- Total des surplus = Σ écarts (montant) > 0 ; Total des manques = |Σ écarts (montant) < 0|
- Écart financier net = surplus − manques ; Capital réel du stock = Total physique
  (inventaire_service.resume_session)

Le fichier est un document de travail : il n'a aucun effet sur l'inventaire enregistré.
"""
from __future__ import annotations

import io
import secrets
from decimal import Decimal

from django.utils import timezone
from openpyxl import Workbook
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Protection, Side
from openpyxl.utils import get_column_letter
from openpyxl.workbook.properties import CalcProperties
from openpyxl.worksheet.datavalidation import DataValidation

from stock.models import Devise, InventaireSession

# Couleurs par défaut du rapport (identiques aux défauts de l'apparence des documents du frontend).
COULEURS_DEFAUT = {
    'primary': '#11116f',
    'title': '#11116f',
    'subtitle': '#1e293b',
    'label': '#475569',
    'value': '#0f172a',
    'table_header_bg': '#11116f',
    'table_header_text': '#ffffff',
    'table_row_even': '#ffffff',
    'table_row_odd': '#f8fafc',
    'table_border': '#cbd5e1',
    'total_row_bg': '#f1f5f9',
    'total_row_text': '#0f172a',
}
COULEUR_SAISIE = 'FFF4CC'       # cellules modifiables (jaune clair)
COULEUR_SAISIE_BORD = 'D4A72C'
COULEUR_POSITIF = '15803D'
COULEUR_NEGATIF = 'B91C1C'

COLONNES = [
    # (titre, largeur)
    ('Article', 36),
    ('Unité', 12),
    ('Stock logiciel', 15),
    ('Stock physique', 15),
    ('Écart', 11),
    ('Dernier PU', 14),
    ('Montant logiciel', 18),
    ('Montant physique', 18),
    ('Écart (montant)', 18),
]
NB_COLONNES = len(COLONNES)
DERNIERE_COLONNE = get_column_letter(NB_COLONNES)

PERIMETRES = {
    InventaireSession.PERIMETRE_COMPLET: 'Catalogue complet',
    InventaireSession.PERIMETRE_EN_STOCK: 'Articles en stock uniquement',
    InventaireSession.PERIMETRE_PARTIEL: "Liste d'articles",
}


def _hex(couleur: str | None, defaut: str) -> str:
    valeur = (couleur or defaut or '').strip().lstrip('#')
    if len(valeur) == 3:
        valeur = ''.join(c * 2 for c in valeur)
    if len(valeur) != 6 or any(c not in '0123456789abcdefABCDEF' for c in valeur):
        valeur = defaut.lstrip('#')
    return valeur.upper()


def _couleurs(entreprise) -> dict[str, str]:
    couleurs = dict(COULEURS_DEFAUT)
    try:
        config = entreprise.get_config_dict() if entreprise else {}
        personnalisees = ((config.get('document_appearance') or {}).get('inventaire') or {}).get('colors') or {}
    except Exception:  # configuration illisible : couleurs par défaut
        personnalisees = {}
    if isinstance(personnalisees, dict):
        couleurs.update({k: v for k, v in personnalisees.items() if isinstance(v, str)})
    return {k: _hex(v, COULEURS_DEFAUT.get(k, '#000000')) for k, v in couleurs.items()}


def _format_monnaie(devise: Devise | None) -> tuple[str, str]:
    """(format montant, format montant signé) avec le symbole de la devise principale."""
    symbole = (devise.symbole or devise.sigle or '').strip() if devise else ''
    suffixe = f' "{symbole}"' if symbole else ''
    simple = f'#,##0.00{suffixe};-#,##0.00{suffixe}'
    signe = f'+#,##0.00{suffixe};-#,##0.00{suffixe};0.00{suffixe}'
    return simple, signe


def _nom_utilisateur(user) -> str:
    if user is None:
        return ''
    return (user.get_full_name() or '').strip() or user.username


def _decimal(valeur) -> float | None:
    """Valeur numérique Excel (float) en conservant les décimales enregistrées."""
    if valeur is None:
        return None
    nombre = Decimal(str(valeur)).normalize()
    return int(nombre) if nombre == nombre.to_integral_value() else float(nombre)


def nom_fichier_export(session: InventaireSession) -> str:
    base = ''.join(c if c.isalnum() or c in '-_' else '_' for c in (session.libelle or '').strip())[:60]
    return f'inventaire_{session.pk}_{base or "rapport"}_{session.date_inventaire:%Y%m%d}.xlsx'


def construire_classeur_inventaire(session: InventaireSession) -> bytes:
    entreprise = session.entreprise
    couleurs = _couleurs(entreprise)
    devise = (
        Devise.objects.filter(entreprise=entreprise, est_principal=True).first()
        or Devise.objects.filter(entreprise=entreprise).first()
    )
    fmt_monnaie, fmt_monnaie_signe = _format_monnaie(devise)
    fmt_qte = 'General'
    fmt_qte_signe = '+General;-General;0'

    lignes = list(
        session.lignes.select_related('article__unite').order_by('article__nom_scientifique', 'article_id')
    )

    wb = Workbook()
    ws = wb.active
    ws.title = 'Inventaire'
    ws.sheet_view.showGridLines = False
    wb.calculation = CalcProperties(fullCalcOnLoad=True)
    wb.properties.title = f"Rapport d'inventaire — {session.libelle}"
    wb.properties.creator = entreprise.nom if entreprise else 'UHAKIKAAPP'

    for index, (_titre, largeur) in enumerate(COLONNES, start=1):
        ws.column_dimensions[get_column_letter(index)].width = largeur

    fin = Side(style='thin', color=couleurs['table_border'])
    bordure = Border(left=fin, right=fin, top=fin, bottom=fin)
    bord_saisie = Side(style='thin', color=COULEUR_SAISIE_BORD)
    bordure_saisie = Border(left=bord_saisie, right=bord_saisie, top=bord_saisie, bottom=bord_saisie)
    remplissage = lambda couleur: PatternFill('solid', start_color=couleur, end_color=couleur)  # noqa: E731

    def fusion(ligne: int, valeur, *, police: Font, alignement='left', hauteur=None):
        ws.merge_cells(start_row=ligne, start_column=1, end_row=ligne, end_column=NB_COLONNES)
        cellule = ws.cell(row=ligne, column=1, value=valeur)
        cellule.font = police
        cellule.alignment = Alignment(horizontal=alignement, vertical='center', wrap_text=True)
        if hauteur:
            ws.row_dimensions[ligne].height = hauteur
        return cellule

    # --- En-tête de l'établissement -------------------------------------------------------
    r = 1
    fusion(r, entreprise.nom if entreprise else '', police=Font(name='Arial', size=14, bold=True, color=couleurs['primary']), hauteur=22)
    r += 1
    coordonnees = ' · '.join(
        x for x in [
            getattr(entreprise, 'adresse', ''),
            f"Tél. : {entreprise.telephone}" if getattr(entreprise, 'telephone', '') else '',
            getattr(entreprise, 'email', ''),
        ] if x
    )
    identite = ' · '.join(
        x for x in [
            f"NIF : {entreprise.nif}" if getattr(entreprise, 'nif', '') else '',
            getattr(entreprise, 'secteur', ''),
            getattr(entreprise, 'pays', ''),
        ] if x
    )
    for texte in (coordonnees, identite):
        if texte:
            fusion(r, texte, police=Font(name='Arial', size=9, color=couleurs['label']))
            r += 1
    if session.succursale_id:
        fusion(r, f"Agence : {session.succursale.nom}", police=Font(name='Arial', size=9, color=couleurs['label']))
        r += 1

    # Filet de séparation sous l'en-tête.
    for col in range(1, NB_COLONNES + 1):
        ws.cell(row=r, column=col).border = Border(bottom=Side(style='medium', color=couleurs['primary']))
    r += 2

    # --- Titre ----------------------------------------------------------------------------
    fusion(r, "RAPPORT D'INVENTAIRE", police=Font(name='Arial', size=18, bold=True, color=couleurs['title']), alignement='center', hauteur=28)
    r += 1
    fusion(r, session.libelle, police=Font(name='Arial', size=11, color=couleurs['subtitle']), alignement='center')
    r += 2

    # --- Informations de la session (deux colonnes libellé / valeur) -----------------------
    date_validation = (
        timezone.localtime(session.date_validation).strftime('%d/%m/%Y %H:%M') if session.date_validation else '—'
    )
    infos = [
        ('Statut', session.get_statut_display()),
        ("Date d'inventaire", session.date_inventaire.strftime('%d/%m/%Y')),
        ('Période', f"Du {session.date_inventaire:%d/%m/%Y} au {session.date_inventaire:%d/%m/%Y}"),
        ('Périmètre', PERIMETRES.get(session.perimetre, session.perimetre)),
        ('Créé par', _nom_utilisateur(session.cree_par) or '—'),
        ('Validé par', _nom_utilisateur(session.valide_par) or '—'),
        ('Validé le', date_validation),
        ('Devise', f"{devise.nom} ({devise.sigle})" if devise else '—'),
    ]
    police_libelle = Font(name='Arial', size=10, bold=True, color=couleurs['label'])
    police_valeur = Font(name='Arial', size=10, color=couleurs['value'])
    # Colonne de gauche : libellé A, valeur B:D ; colonne de droite : libellé F:G, valeur H:I.
    positions = ((1, 1, 2, 4), (6, 7, 8, 9))
    for i in range(0, len(infos), 2):
        for (lib_debut, lib_fin, val_debut, val_fin), (libelle, valeur) in zip(positions, infos[i:i + 2]):
            if lib_fin > lib_debut:
                ws.merge_cells(start_row=r, start_column=lib_debut, end_row=r, end_column=lib_fin)
            ws.cell(row=r, column=lib_debut, value=f'{libelle} :').font = police_libelle
            ws.merge_cells(start_row=r, start_column=val_debut, end_row=r, end_column=val_fin)
            cellule = ws.cell(row=r, column=val_debut, value=valeur)
            cellule.font = police_valeur
        r += 1
    r += 1

    # --- Positions du tableau (connues avant les indicateurs, qui y font référence) -------
    ligne_indicateurs = r
    ligne_aide = ligne_indicateurs + 3
    ligne_entete = ligne_aide + 2
    premiere = ligne_entete + 1
    derniere = premiere + max(len(lignes), 1) - 1
    plage = lambda col: f'${col}${premiere}:${col}${derniere}'  # noqa: E731

    # --- Indicateurs (cartes) : recalculés à partir du stock physique saisi --------------
    indicateurs = [
        ('Comptées', f'=COUNT({plage("D")})'),
        ('Restantes', f'=ROWS({plage("C")})-COUNT({plage("D")})'),
        ('Avec écart', f'=COUNTIF({plage("E")},">0")+COUNTIF({plage("E")},"<0")'),
        ('Surplus', f'=COUNTIF({plage("E")},">0")'),
        ('Manquants', f'=COUNTIF({plage("E")},"<0")'),
        ('Conformes', f'=COUNTIF({plage("E")},0)'),
    ]
    if not lignes:
        indicateurs = [(libelle, 0) for libelle, _formule in indicateurs]
    fond_carte = remplissage(couleurs['total_row_bg'])
    for index, (libelle, formule) in enumerate(indicateurs):
        col = 3 + index  # colonnes C…H, de largeurs voisines
        haut = ws.cell(row=ligne_indicateurs, column=col, value=libelle)
        haut.font = Font(name='Arial', size=9, bold=True, color=couleurs['label'])
        bas = ws.cell(row=ligne_indicateurs + 1, column=col, value=formule)
        bas.font = Font(name='Arial', size=14, bold=True, color=couleurs['value'])
        bas.number_format = '0'
        for cellule in (haut, bas):
            cellule.fill = fond_carte
            cellule.border = bordure
            cellule.alignment = Alignment(horizontal='center', vertical='center')
    ws.row_dimensions[ligne_indicateurs + 1].height = 24

    # --- Aide à la saisie ------------------------------------------------------------------
    aide = fusion(
        ligne_aide,
        "Cellules jaunes modifiables : « Stock physique » et « Dernier PU ». Les écarts, montants, "
        "indicateurs et totaux se recalculent automatiquement. Ce fichier est un document de travail : "
        "il ne modifie pas l'inventaire enregistré dans l'application.",
        police=Font(name='Arial', size=9, italic=True, color=couleurs['label']),
        hauteur=30,
    )
    aide.fill = remplissage(COULEUR_SAISIE)

    # --- Tableau des articles ----------------------------------------------------------------
    for index, (titre, _largeur) in enumerate(COLONNES, start=1):
        cellule = ws.cell(row=ligne_entete, column=index, value=titre)
        cellule.font = Font(name='Arial', size=10, bold=True, color=couleurs['table_header_text'])
        cellule.fill = remplissage(couleurs['table_header_bg'])
        cellule.border = bordure
        cellule.alignment = Alignment(horizontal='left' if index <= 2 else 'right', vertical='center', wrap_text=True)
    ws.row_dimensions[ligne_entete].height = 30

    deverrouillee = Protection(locked=False)
    for position, ligne in enumerate(lignes):
        n = premiere + position
        article = ligne.article
        fond = remplissage(couleurs['table_row_even'] if position % 2 == 0 else couleurs['table_row_odd'])
        valeurs = [
            article.nom_commercial or article.nom_scientifique or article.article_id,
            article.unite.libelle if article.unite_id else '',
            _decimal(ligne.stock_theorique),
            _decimal(ligne.stock_physique),
            # Arrondis à 5 décimales comme le backend (quantize 0.00001) : pas de résidu flottant
            # qui ferait compter un écart nul comme surplus ou manque.
            f'=IF(D{n}="","",ROUND(D{n}-C{n},5))',
            _decimal(ligne.dernier_prix_unitaire),
            f'=ROUND(C{n}*F{n},5)',
            f'=IF(D{n}="","",ROUND(D{n}*F{n},5))',
            f'=IF(D{n}="","",ROUND(H{n}-G{n},5))',
        ]
        formats = ['@', '@', fmt_qte, fmt_qte, fmt_qte_signe, fmt_monnaie, fmt_monnaie, fmt_monnaie, fmt_monnaie_signe]
        for col, (valeur, format_nombre) in enumerate(zip(valeurs, formats), start=1):
            cellule = ws.cell(row=n, column=col, value=valeur)
            cellule.number_format = format_nombre
            cellule.font = Font(name='Arial', size=10, color=couleurs['value'])
            cellule.alignment = Alignment(horizontal='left' if col <= 2 else 'right', vertical='center')
            if col in (4, 6):  # Stock physique, Dernier PU : seules cellules modifiables
                cellule.protection = deverrouillee
                cellule.fill = remplissage(COULEUR_SAISIE)
                cellule.border = bordure_saisie
                cellule.font = Font(name='Arial', size=10, bold=True, color=couleurs['value'])
            else:
                cellule.fill = fond
                cellule.border = bordure

    # Contrôle de saisie : nombres positifs ou nuls uniquement.
    if lignes:
        for col, message in (('D', 'Saisissez une quantité positive ou nulle.'), ('F', 'Saisissez un prix positif ou nul.')):
            validation = DataValidation(
                type='decimal', operator='greaterThanOrEqual', formula1='0', allow_blank=True,
                showErrorMessage=True, errorTitle='Valeur invalide', error=message,
            )
            validation.add(f'{col}{premiere}:{col}{derniere}')
            ws.add_data_validation(validation)
        # Écarts en couleur : vert pour un surplus, rouge pour un manque.
        for col in ('E', 'I'):
            zone = f'{col}{premiere}:{col}{derniere}'
            ws.conditional_formatting.add(zone, CellIsRule(operator='greaterThan', formula=['0'], font=Font(color=COULEUR_POSITIF, bold=True)))
            ws.conditional_formatting.add(zone, CellIsRule(operator='lessThan', formula=['0'], font=Font(color=COULEUR_NEGATIF, bold=True)))

    # --- Ligne de totaux -------------------------------------------------------------------
    ligne_total = derniere + 1
    totaux = {
        1: 'TOTAL',
        7: f'=SUM({plage("G")})' if lignes else 0,
        8: f'=SUM({plage("H")})' if lignes else 0,
        9: f'=SUM({plage("I")})' if lignes else 0,
    }
    for col in range(1, NB_COLONNES + 1):
        cellule = ws.cell(row=ligne_total, column=col, value=totaux.get(col))
        cellule.fill = remplissage(couleurs['total_row_bg'])
        cellule.border = bordure
        cellule.font = Font(name='Arial', size=10, bold=True, color=couleurs['total_row_text'])
        cellule.alignment = Alignment(horizontal='left' if col <= 2 else 'right', vertical='center')
        cellule.number_format = fmt_monnaie_signe if col == 9 else fmt_monnaie
    if lignes:
        ws.conditional_formatting.add(f'I{ligne_total}', CellIsRule(operator='greaterThan', formula=['0'], font=Font(color=COULEUR_POSITIF, bold=True)))
        ws.conditional_formatting.add(f'I{ligne_total}', CellIsRule(operator='lessThan', formula=['0'], font=Font(color=COULEUR_NEGATIF, bold=True)))

    # --- Pied financier (libellé / valeur) -------------------------------------------------
    r = ligne_total + 2
    ref = {}
    pied = [
        ('total_logiciel', 'Total logiciel', lambda: f'=G{ligne_total}', fmt_monnaie, False),
        ('total_physique', 'Total physique', lambda: f'=H{ligne_total}', fmt_monnaie, False),
        ('ecart_financier', 'Écart financier', lambda: f"=ROUND({ref['total_physique']}-{ref['total_logiciel']},5)", fmt_monnaie_signe, False),
        ('total_surplus', 'Total des surplus', lambda: f'=SUMIF({plage("I")},">0")' if lignes else 0, fmt_monnaie_signe, False),
        ('total_manques', 'Total des manques', lambda: f'=ABS(SUMIF({plage("I")},"<0"))' if lignes else 0, fmt_monnaie, False),
        ('ecart_net', 'Écart financier net', lambda: f"=ROUND({ref['total_surplus']}-{ref['total_manques']},5)", fmt_monnaie_signe, False),
        ('capital_reel', 'Capital réel du stock', lambda: f"={ref['total_physique']}", fmt_monnaie, True),
    ]
    for cle, libelle, formule, format_nombre, en_avant in pied:
        ws.merge_cells(start_row=r, start_column=6, end_row=r, end_column=8)
        etiquette = ws.cell(row=r, column=6, value=libelle)
        valeur = ws.cell(row=r, column=9, value=formule())
        ref[cle] = f'I{r}'
        fond = remplissage(couleurs['primary'] if en_avant else couleurs['total_row_bg'])
        texte = couleurs['table_header_text'] if en_avant else couleurs['label']
        for col in (6, 7, 8, 9):
            cellule = ws.cell(row=r, column=col)
            cellule.fill = fond
            cellule.border = bordure
        etiquette.font = Font(name='Arial', size=11 if en_avant else 10, bold=True, color=texte)
        etiquette.alignment = Alignment(horizontal='left', vertical='center', indent=1)
        valeur.font = Font(
            name='Arial', size=13 if en_avant else 10, bold=True,
            color=couleurs['table_header_text'] if en_avant else couleurs['value'],
        )
        valeur.alignment = Alignment(horizontal='right', vertical='center')
        valeur.number_format = format_nombre
        if en_avant:
            ws.row_dimensions[r].height = 24
        elif cle in ('ecart_financier', 'ecart_net', 'total_surplus'):
            ws.conditional_formatting.add(valeur.coordinate, CellIsRule(operator='greaterThan', formula=['0'], font=Font(color=COULEUR_POSITIF, bold=True)))
            ws.conditional_formatting.add(valeur.coordinate, CellIsRule(operator='lessThan', formula=['0'], font=Font(color=COULEUR_NEGATIF, bold=True)))
        r += 1

    r += 1
    fusion(
        r,
        f"Exporté le {timezone.localtime():%d/%m/%Y à %H:%M} — valeurs de l'inventaire validé ; "
        "les cellules jaunes peuvent être ajustées pour simulation.",
        police=Font(name='Arial', size=8, italic=True, color=couleurs['label']),
    )

    # --- Mise en page et protection ----------------------------------------------------------
    ws.freeze_panes = ws.cell(row=premiere, column=1)
    ws.print_title_rows = f'{ligne_entete}:{ligne_entete}'
    ws.print_area = f'A1:{DERNIERE_COLONNE}{r}'
    ws.page_setup.orientation = 'landscape'
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_margins.left = ws.page_margins.right = 0.4
    ws.oddFooter.center.text = 'Page &P / &N'

    # Mot de passe aléatoire non conservé : les cellules calculées ne peuvent pas être modifiées.
    ws.protection.password = secrets.token_urlsafe(18)
    ws.protection.sheet = True
    ws.protection.selectLockedCells = False
    ws.protection.selectUnlockedCells = False
    ws.protection.formatColumns = False
    ws.protection.formatRows = False
    ws.protection.sort = True
    ws.protection.autoFilter = True

    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()
