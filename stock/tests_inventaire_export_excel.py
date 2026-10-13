"""Export Excel d'un inventaire validé : données réelles, formules, protection, permissions."""
import io
from decimal import Decimal

from django.contrib.auth import get_user_model
from django.test import override_settings
from django.utils import timezone
from openpyxl import load_workbook
from rest_framework.test import APITestCase

from stock.models import (
    Article,
    Devise,
    Entree,
    Entreprise,
    InventaireLigne,
    InventaireSession,
    LigneEntree,
    SousTypeArticle,
    Stock,
    TypeArticle,
    Unite,
)
from users.models import Membership

D = Decimal
XLSX = 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'


def _entreprise(nom, email):
    return Entreprise.objects.create(
        nom=nom, secteur='Commerce', pays='CD', adresse='Av. du Marché 12', telephone='+243 81 000 0000',
        email=email, nif=f'NIF-{nom}', responsable='Resp',
    )


@override_settings(ALLOWED_HOSTS=['testserver', 'localhost', '127.0.0.1'])
class InventaireExportExcelTests(APITestCase):
    def setUp(self):
        self.entreprise = _entreprise('Boutique Excel', 'excel@example.com')
        User = get_user_model()
        self.admin = User.objects.create_user(username='admin_xls', email='a-xls@example.com', password='x-secret-123',
                                              first_name='Awa', last_name='Kabeya')
        Membership.objects.create(user=self.admin, entreprise=self.entreprise, role='admin', is_active=True)
        self.agent = User.objects.create_user(username='agent_xls', email='g-xls@example.com', password='x-secret-123')
        Membership.objects.create(user=self.agent, entreprise=self.entreprise, role='user', is_active=True)
        self.client.force_authenticate(user=self.admin)
        Devise.objects.create(sigle='USD', nom='Dollar', symbole='$', est_principal=True, entreprise=self.entreprise)
        unite = Unite.objects.create(libelle='Pièce', entreprise=self.entreprise)
        type_art = TypeArticle.objects.create(libelle='Produit', entreprise=self.entreprise)
        sous = SousTypeArticle.objects.create(libelle='Divers', type_article=type_art, entreprise=self.entreprise)
        entree = Entree.objects.create(libele='Appro', entreprise=self.entreprise)
        self.articles = {}
        for nom, stock, pu in (('Adroid', '46', '0.25'), ('Amara', '29', '0.40'), ('Bic', '10', '1.50')):
            article = Article.objects.create(
                nom_scientifique=nom, nom_commercial=nom, sous_type_article=sous, unite=unite,
                entreprise=self.entreprise,
            )
            Stock.objects.create(article=article, Qte=D(stock), seuilAlert=D('0'))
            LigneEntree.objects.create(
                article=article, entree=entree, quantite=D(stock), quantite_restante=D(stock),
                prix_unitaire=D(pu), prix_vente=D('5'),
            )
            self.articles[nom] = article

    def _inventaire_valide(self):
        create = self.client.post('/api/inventaires/', {
            'libelle': 'Inventaire de fin de mois', 'date_inventaire': timezone.localdate().isoformat(),
            'perimetre': 'EN_STOCK', 'demarrer': True,
        }, format='json')
        self.assertEqual(create.status_code, 201, create.content)
        session_id = create.json()['id']
        comptes = {'Adroid': '44', 'Amara': '30', 'Bic': '10'}  # manque, surplus, conforme
        bulk = self.client.post(f'/api/inventaires/{session_id}/lignes/bulk/', {'lignes': [
            {'article_id': self.articles[nom].article_id, 'stock_physique': qte} for nom, qte in comptes.items()
        ]}, format='json')
        self.assertEqual(bulk.status_code, 200, bulk.content)
        valider = self.client.post(f'/api/inventaires/{session_id}/valider/', {}, format='json')
        self.assertEqual(valider.status_code, 200, valider.content)
        return InventaireSession.objects.get(pk=session_id)

    def _export(self, session_id, expect=200):
        response = self.client.get(f'/api/inventaires/{session_id}/export-excel/')
        self.assertEqual(response.status_code, expect, getattr(response, 'content', b'')[:300])
        return response

    @staticmethod
    def _classeur(response):
        return load_workbook(io.BytesIO(response.content))

    @staticmethod
    def _trouver(ws, valeur):
        for row in ws.iter_rows():
            for cellule in row:
                if cellule.value == valeur:
                    return cellule
        raise AssertionError(f'{valeur!r} introuvable')

    def test_export_reprend_le_rapport_avec_les_donnees_reelles(self):
        session = self._inventaire_valide()
        response = self._export(session.pk)
        self.assertEqual(response['Content-Type'], XLSX)
        self.assertIn(f'inventaire_{session.pk}_', response['Content-Disposition'])
        ws = self._classeur(response)['Inventaire']

        self.assertEqual(ws['A1'].value, 'Boutique Excel')
        self.assertTrue(any('NIF : NIF-Boutique Excel' in str(c.value) for c in ws['A'] if c.value))
        self._trouver(ws, "RAPPORT D'INVENTAIRE")
        self._trouver(ws, 'Inventaire de fin de mois')
        def info(libelle):
            cellule = self._trouver(ws, libelle)
            return next(
                ws.cell(row=cellule.row, column=c).value
                for c in range(cellule.column + 1, 10)
                if ws.cell(row=cellule.row, column=c).value not in (None, '')
            )

        self.assertEqual(info('Statut :'), 'Validé')
        self.assertEqual(info('Validé par :'), 'Awa Kabeya')
        self.assertEqual(info('Devise :'), 'Dollar (USD)')
        self._trouver(ws, 'Période :')
        for libelle in ('Comptées', 'Restantes', 'Avec écart', 'Surplus', 'Manquants', 'Conformes'):
            self._trouver(ws, libelle)

        entete = self._trouver(ws, 'Article')
        titres = [ws.cell(row=entete.row, column=c).value for c in range(1, 10)]
        self.assertEqual(titres, [
            'Article', 'Unité', 'Stock logiciel', 'Stock physique', 'Écart', 'Dernier PU',
            'Montant logiciel', 'Montant physique', 'Écart (montant)',
        ])
        premiere = entete.row + 1
        lignes = {ws.cell(row=premiere + i, column=1).value: premiere + i for i in range(3)}
        n = lignes['Adroid']
        self.assertEqual(ws[f'B{n}'].value, 'Pièce')
        self.assertEqual((ws[f'C{n}'].value, ws[f'D{n}'].value, ws[f'F{n}'].value), (46, 44, 0.25))
        self.assertEqual(ws[f'E{n}'].value, f'=IF(D{n}="","",ROUND(D{n}-C{n},5))')
        self.assertEqual(ws[f'G{n}'].value, f'=ROUND(C{n}*F{n},5)')
        self.assertEqual(ws[f'H{n}'].value, f'=IF(D{n}="","",ROUND(D{n}*F{n},5))')
        self.assertEqual(ws[f'I{n}'].value, f'=IF(D{n}="","",ROUND(H{n}-G{n},5))')
        self.assertEqual((ws[f'C{lignes["Amara"]}'].value, ws[f'D{lignes["Amara"]}'].value), (29, 30))
        self.assertIn('"$"', ws[f'F{n}'].number_format)
        self.assertTrue(ws[f'I{n}'].number_format.startswith('+'))

        derniere = premiere + 2
        total = self._trouver(ws, 'TOTAL')
        self.assertEqual(total.row, derniere + 1)
        self.assertEqual(ws[f'G{total.row}'].value, f'=SUM($G${premiere}:$G${derniere})')
        self.assertEqual(ws[f'I{total.row}'].value, f'=SUM($I${premiere}:$I${derniere})')

        def valeur_pied(libelle):
            return ws.cell(row=self._trouver(ws, libelle).row, column=9).value

        self.assertEqual(valeur_pied('Total logiciel'), f'=G{total.row}')
        self.assertEqual(valeur_pied('Total physique'), f'=H{total.row}')
        self.assertEqual(valeur_pied('Total des surplus'), f'=SUMIF($I${premiere}:$I${derniere},">0")')
        self.assertEqual(valeur_pied('Total des manques'), f'=ABS(SUMIF($I${premiere}:$I${derniere},"<0"))')
        ligne_surplus = self._trouver(ws, 'Total des surplus').row
        self.assertEqual(valeur_pied('Écart financier net'), f'=ROUND(I{ligne_surplus}-I{ligne_surplus + 1},5)')
        ligne_physique = self._trouver(ws, 'Total physique').row
        self.assertEqual(valeur_pied('Capital réel du stock'), f'=I{ligne_physique}')
        self.assertEqual(valeur_pied('Écart financier'), f'=ROUND(I{ligne_physique}-I{ligne_physique - 1},5)')

    def test_seuls_stock_physique_et_dernier_pu_sont_modifiables(self):
        session = self._inventaire_valide()
        ws = self._classeur(self._export(session.pk))['Inventaire']
        self.assertTrue(ws.protection.sheet)
        self.assertTrue(ws.protection.password)
        entete = self._trouver(ws, 'Article')
        modifiables = set()
        for row in ws.iter_rows():
            for cellule in row:
                if not cellule.protection.locked:
                    modifiables.add(cellule.coordinate)
        attendues = {f'{col}{entete.row + i}' for col in ('D', 'F') for i in range(1, 4)}
        self.assertEqual(modifiables, attendues)
        self.assertEqual(ws[f'D{entete.row + 1}'].fill.start_color.rgb[-6:], 'FFF4CC')
        validations = {str(v.sqref): v.formula1 for v in ws.data_validations.dataValidation}
        self.assertEqual(validations, {f'D{entete.row + 1}:D{entete.row + 3}': '0', f'F{entete.row + 1}:F{entete.row + 3}': '0'})

    def test_couleurs_de_l_apparence_du_rapport(self):
        config = self.entreprise.get_config_dict()
        config.setdefault('document_appearance', {})['inventaire'] = {'colors': {'table_header_bg': '#7a1f2b'}}
        self.entreprise.set_config_dict(config)
        self.entreprise.save(update_fields=['config'])
        session = self._inventaire_valide()
        ws = self._classeur(self._export(session.pk))['Inventaire']
        self.assertEqual(self._trouver(ws, 'Article').fill.start_color.rgb[-6:], '7A1F2B')

    def test_export_sans_effet_sur_l_inventaire(self):
        session = self._inventaire_valide()
        avant = list(InventaireLigne.objects.filter(session=session).values_list('stock_physique', 'dernier_prix_unitaire'))
        self._export(session.pk)
        self.assertEqual(
            list(InventaireLigne.objects.filter(session=session).values_list('stock_physique', 'dernier_prix_unitaire')),
            avant,
        )
        self.assertEqual(InventaireSession.objects.get(pk=session.pk).statut, InventaireSession.STATUT_VALIDE)

    def test_reserve_aux_inventaires_valides(self):
        create = self.client.post('/api/inventaires/', {
            'libelle': 'En cours', 'date_inventaire': timezone.localdate().isoformat(),
            'perimetre': 'EN_STOCK', 'demarrer': True,
        }, format='json')
        response = self._export(create.json()['id'], expect=400)
        self.assertIn('validé', response.json()['detail'])

    def test_agent_de_l_entreprise_peut_exporter(self):
        session = self._inventaire_valide()
        self.client.force_authenticate(user=self.agent)
        self._export(session.pk)

    def test_isolation_entre_entreprises(self):
        session = self._inventaire_valide()
        autre = _entreprise('Autre', 'autre-xls@example.com')
        intrus = get_user_model().objects.create_user(username='intrus_xls', email='i-xls@example.com', password='x-secret-123')
        Membership.objects.create(user=intrus, entreprise=autre, role='admin', is_active=True)
        self.client.force_authenticate(user=intrus)
        self._export(session.pk, expect=404)
        self.client.force_authenticate(user=None)
        self.assertIn(self.client.get(f'/api/inventaires/{session.pk}/export-excel/').status_code, (401, 403))
